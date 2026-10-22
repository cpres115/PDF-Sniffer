# Author:  Corey Preston
# Company: Fischer Connectors, Inc
# Date:    October 2026
import csv
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

from pdf_sniffer import models, namecheck, selftest
from pdf_sniffer.config import ScanConfig
from pdf_sniffer.scanner import scan_file


def result(path, category, labels=(), status="OK"):
    return {"path": path, "category": category, "status": status, "notes": [], "suspects": [],
            "hits": [{"label": label, "eccn": "", "page": 1, "method": "text", "confidence": "high",
                      "evidence": ""} for label in labels]}


class VerdictTests(unittest.TestCase):
    def test_decision_table(self):
        v = namecheck.verdict
        self.assertEqual(v("", ""), (namecheck.MATCH, ""))
        self.assertEqual(v("EAR", "EAR"), (namecheck.MATCH, ""))
        self.assertEqual(v("ITAR", "ITAR"), (namecheck.MATCH, ""))
        for name in ("", "EAR", "ITAR"):
            self.assertEqual(v(name, None), (namecheck.NOT_COMPARED, ""))
        # logo stricter than the name -> High
        self.assertEqual(v("", "EAR"), ("Name has no ITAR/EAR tag, but the file carries an EAR logo", "High"))
        self.assertEqual(v("", "ITAR"), ("Name has no ITAR/EAR tag, but the file carries an ITAR logo", "High"))
        self.assertEqual(v("EAR", "ITAR"), ("Name says EAR, logo says ITAR", "High"))
        # name stricter than the logo -> Medium
        self.assertEqual(v("ITAR", "EAR"), ("Name says ITAR, logo says EAR", "Medium"))
        self.assertEqual(v("EAR", ""), ("Name says EAR, no ITAR/EAR logo found", "Medium"))
        self.assertEqual(v("ITAR", ""), ("Name says ITAR, no ITAR/EAR logo found", "Medium"))

    def test_logo_classification(self):
        lc = namecheck.logo_classification
        self.assertEqual(lc(result("a.pdf", models.NON_CONTROLLED)), "")
        self.assertEqual(lc(result("a.pdf", models.EAR_NO_ECCN, ["EAR"])), "EAR")
        self.assertEqual(lc(result("a.pdf", models.ITAR, ["ITAR", "EAR"])), "ITAR")   # both: the stricter wins
        self.assertIsNone(lc(result("a.pdf", models.NEEDS_REVIEW)))
        self.assertIsNone(lc(result("a.pdf", models.UNREADABLE)))


class CheckTests(unittest.TestCase):
    def test_check_marks_mismatches_and_counts(self):
        rs = [result("/lib/A-1_A-ITAR.pdf", models.ITAR, ["ITAR"]),                    # agree
              result("/lib/A-2_B-EAR.pdf", models.ITAR, ["ITAR"]),                     # High
              result("/lib/A-3_A.pdf", "EAR Controlled - ECCN 9A610", ["EAR"]),        # High
              result("/lib/A-4_A-EAR99.pdf", models.EAR_NO_ECCN, ["EAR"]),             # agree (EAR99 = EAR)
              result("/lib/A-5_A-ITAR.pdf", models.EAR_NO_ECCN, ["EAR"]),              # Medium
              result("/lib/A-6_A-EAR.pdf", models.NON_CONTROLLED),                     # Medium
              result("/lib/A-7_A.pdf", models.NON_CONTROLLED),                         # agree
              result("/lib/A-8_A-EAR.pdf", models.NEEDS_REVIEW, status="REVIEW"),      # cannot compare
              result("/lib/A-9_A-ITAR.pdf", models.UNREADABLE, status="ERROR")]        # cannot compare
        out = namecheck.check(rs)
        self.assertEqual(out.stats, {"compared": 7, "match": 3, "mismatch": 4, "high": 2, "medium": 2,
                                     "not_compared": 2})
        self.assertEqual([row["Priority"] for row in out.rows], ["High", "High", "Medium", "Medium"])
        self.assertEqual({row["File Name"] for row in out.rows[:2]}, {"A-2_B-EAR.pdf", "A-3_A.pdf"})
        # mismatches are flagged for review and get a note; matches and uncomparable files are left alone
        by_name = {os.path.basename(r["path"]): r for r in rs}
        self.assertEqual(by_name["A-2_B-EAR.pdf"]["status"], "REVIEW")
        self.assertIn("Name vs logo: Name says EAR, logo says ITAR", by_name["A-2_B-EAR.pdf"]["notes"])
        self.assertEqual(by_name["A-1_A-ITAR.pdf"]["status"], "OK")
        self.assertEqual(by_name["A-9_A-ITAR.pdf"]["status"], "ERROR")
        self.assertEqual(by_name["A-7_A.pdf"]["name_vs_logo"], namecheck.MATCH)
        self.assertEqual(by_name["A-8_A-EAR.pdf"]["name_vs_logo"], namecheck.NOT_COMPARED)


class EndToEndTests(unittest.TestCase):
    """Real sample PDFs under file names that carry (or lack) a classification tag."""
    NAMES = {  # sample PDF -> {library file name: expected verdict}
        "02_document_itar_footer_vector.pdf": {"TST-000-101_A-ITAR.pdf": "Match",
                                              "TST-000-102_A-EAR.pdf": "Name says EAR, logo says ITAR",
                                              "TST-000-103_A.pdf": "Name has no ITAR/EAR tag, but the file carries an ITAR logo"},
        "04_document_ear_eccn_9A610x.pdf": {"TST-000-104_A-EAR.pdf": "Match", "TST-000-107_A-EAR99.pdf": "Match",
                                            "TST-000-105_A-ITAR.pdf": "Name says ITAR, logo says EAR",
                                            "TST-000-106_A.pdf": "Name has no ITAR/EAR tag, but the file carries an EAR logo"},
        "01_document_plain.pdf": {"TST-000-108_A.pdf": "Match",
                                  "TST-000-109_A-EAR.pdf": "Name says EAR, no ITAR/EAR logo found"},
        "19_mwi_advisory_table_only.pdf": {"MWI-000-110_A.pdf": "Match"},   # explanatory logos must not count
        "11_document_itar_and_ear.pdf": {"TST-000-111_A-EAR.pdf": "Name says EAR, logo says ITAR"},
        "14_corrupt.pdf": {"TST-000-112_A-ITAR.pdf": "Could not compare"},
        "18_document_red_ring_other_text.pdf": {"TST-000-113_A-EAR.pdf": "Could not compare"},
    }

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="pdf_sniffer_namecheck_")
        samples = os.path.join(cls.tmp, "samples")
        selftest.build_samples(samples)
        cls.library = os.path.join(cls.tmp, "library")
        os.makedirs(cls.library)
        for sample, names in cls.NAMES.items():
            for name in names:
                shutil.copy(os.path.join(samples, sample), os.path.join(cls.library, name))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_verdict_for_every_file(self):
        expected = {name: verdict for names in self.NAMES.values() for name, verdict in names.items()}
        results = [scan_file(os.path.join(self.library, name), ScanConfig()).to_dict() for name in expected]
        out = namecheck.check(results)
        got = {os.path.basename(r["path"]): r["name_vs_logo"] for r in results}
        self.assertEqual(got, expected)
        self.assertEqual((out.stats["compared"], out.stats["match"], out.stats["high"], out.stats["medium"],
                          out.stats["not_compared"]), (11, 5, 4, 2, 2))

    def test_command_line_writes_the_reports(self):
        out_dir = os.path.join(self.tmp, "out")
        proc = subprocess.run([sys.executable, "-m", "pdf_sniffer", self.library, "--out", out_dir, "--workers", "1"],
                              capture_output=True, text=True, cwd=os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertIn("NAME vs LOGO CHECK", proc.stdout)
        run = os.path.join(out_dir, os.listdir(out_dir)[0])
        with open(os.path.join(run, "name_vs_logo.csv"), encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
        self.assertEqual(len(rows), 6)
        self.assertEqual({r["Priority"] for r in rows[:4]}, {"High"})
        self.assertEqual({r["Priority"] for r in rows[4:]}, {"Medium"})
        with open(os.path.join(run, "classification_report.csv"), encoding="utf-8-sig") as fh:
            main = {r["File Name"]: r for r in csv.DictReader(fh)}
        self.assertEqual(main["TST-000-102_A-EAR.pdf"]["Status"], "REVIEW")
        self.assertEqual(main["TST-000-102_A-EAR.pdf"]["Classification In Name"], "EAR")
        self.assertEqual(main["TST-000-103_A.pdf"]["Classification In Name"], "(none)")
        self.assertEqual(main["TST-000-101_A-ITAR.pdf"]["Name vs Logo"], "Match")
        self.assertEqual(main["TST-000-101_A-ITAR.pdf"]["Status"], "OK")
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(run, "classification_report.xlsx"))
        self.assertIn("Name vs Logo", wb.sheetnames)
        self.assertEqual(wb["Name vs Logo"].max_row - 1, 6)


if __name__ == "__main__":
    unittest.main()
