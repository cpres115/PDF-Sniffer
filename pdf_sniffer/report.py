"""CSV / Excel / text reports, plus the optional copy-into-category-folders step."""
import csv
import os
import shutil
import time
from collections import Counter, OrderedDict
from pathlib import Path

from . import models, revisions

COLUMNS = ["Category", "Status", "File Name", "Drawing Number", "Revision", "Folder", "ECCN (as written)", "Logo(s) Found",
           "Logo Pages", "Detection", "Pages", "Pages Scanned", "Size (KB)", "Last Modified",
           "Notes", "Evidence Image(s)", "Full Path"]


def category_order(cat):
    if cat == models.NON_CONTROLLED:
        return (0, cat)
    if cat == models.ITAR:
        return (1, cat)
    if cat == models.EAR_NO_ECCN:
        return (2, cat)
    if cat.startswith(models.EAR_PREFIX):
        return (3, cat)
    if cat == models.NEEDS_REVIEW:
        return (4, cat)
    return (5, cat)


def _row(r):
    hits = r["hits"]
    p = Path(r["path"])
    eccns = sorted({h["eccn"] for h in hits if h["eccn"]})
    labels = Counter(h["label"] for h in hits)
    methods = sorted({f'{h["method"]} ({h["confidence"]})' for h in hits})
    pages = sorted({h["page"] for h in hits + r["suspects"]})
    notes = list(r["notes"])
    notes += [f'p{s["page"]}: {s["detail"]}' for s in r["suspects"] if s["detail"]]
    notes += [f"Info: {i}" for i in r.get("info", [])]
    evidence = [h["evidence"] for h in hits + r["suspects"] if h["evidence"]]
    mtime = time.strftime("%Y-%m-%d %H:%M", time.localtime(r["mtime"])) if r["mtime"] else ""
    return OrderedDict([
        ("Category", r["category"]), ("Status", r["status"]), ("File Name", p.name),
        ("Drawing Number", r.get("drawing_number", "")), ("Revision", r.get("revision", "")),
        ("Folder", str(p.parent)), ("ECCN (as written)", ", ".join(eccns)),
        ("Logo(s) Found", ", ".join(f"{k} x{v}" for k, v in sorted(labels.items()))),
        ("Logo Pages", ", ".join(map(str, pages))), ("Detection", "; ".join(methods)),
        ("Pages", r["pages"]), ("Pages Scanned", r["pages_scanned"]),
        ("Size (KB)", round(r["size"] / 1024, 1)), ("Last Modified", mtime),
        ("Notes", " | ".join(notes)), ("Evidence Image(s)", "; ".join(evidence)),
        ("Full Path", r["path"]),
    ])


def summarize(results):
    return sorted(Counter(r["category"] for r in results).items(), key=lambda kv: category_order(kv[0]))


def revision_summary_lines(stats):
    return ["REVISION CHECK (drawing numbers read from file names)", "-" * 60,
            f"  Drawing numbers present in more than one revision: {stats['multi_revision']}",
            f"  Older revisions that should probably be removed:   {stats['older']}",
            f"  Extra copies of the same revision:                 {stats['extra_copies']}",
            f"  Drawings whose revisions cannot be ordered:        {stats['unordered']}",
            f"  File names with no recognisable number/revision:   {stats['unparsed']} of {stats['files']}"]


def summary_text(results, folders, skipped, seconds, run_dir, rev_stats=None):
    lines = ["PDF SNIFFER - CLASSIFICATION SUMMARY", "=" * 60,
             "Folders scanned:"] + [f"  {f}" for f in folders]
    lines += ["", f"PDF files scanned: {len(results)}   (elapsed {seconds / 60:.1f} min)", ""]
    width = max([len(c) for c, _ in summarize(results)] + [10])
    for cat, n in summarize(results):
        lines.append(f"  {cat.ljust(width)}  {n:>7}")
    st = Counter(r["status"] for r in results)
    lines += ["", f"  Files flagged REVIEW: {st.get(models.STATUS_REVIEW, 0)}    "
                  f"Files that could not be read: {st.get(models.STATUS_ERROR, 0)}"]
    if skipped:
        other = ", ".join(f"{n} x {ext}" for ext, n in skipped.most_common(8))
        lines += ["", f"Non-PDF files ignored: {sum(skipped.values())} ({other})"]
    if rev_stats is not None:
        lines += [""] + revision_summary_lines(rev_stats)
    lines += ["", f"Reports saved in: {run_dir}"]
    return "\n".join(lines)


def write_csv(results, path):
    rows = [_row(r) for r in sorted(results, key=lambda r: (category_order(r["category"]), r["path"].lower()))]
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:  # BOM so Excel reads UTF-8
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)


REV_FILLS = {revisions.STATUS_NEWEST: "E2F0D9", revisions.STATUS_OLDER: "F8CBAD",
             revisions.STATUS_COPY: "FFE699", revisions.STATUS_UNORDERED: "D9D9D9"}


def write_revision_csv(rev_rows, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=revisions.REVISION_COLUMNS)
        w.writeheader()
        w.writerows(rev_rows)


def _add_revision_sheet(wb, rev_rows):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet("Revision Check")
    ws.append(revisions.REVISION_COLUMNS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    for row in rev_rows:
        ws.append([row[k] for k in revisions.REVISION_COLUMNS])
        ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=REV_FILLS[row["Status"]])
    if not rev_rows:
        ws.append(["No drawing number appears more than once."])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"Drawing Number": 20, "Revision": 9, "Status": 34, "File Name": 40, "Folder": 44,
              "Last Modified": 17, "Size (KB)": 10, "Revisions On Server": 20, "Export Control Group": 34,
              "Notes": 44, "Name Read By": 28, "Full Path": 60}
    for i, col in enumerate(revisions.REVISION_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[col]
    for row in ws.iter_rows(min_row=2):
        row[9].alignment = Alignment(wrap_text=True, vertical="top")


def _add_revision_summary(ws, stats):
    from openpyxl.styles import Font
    ws.append([])
    ws.append(["Revision check (from file names)"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for label, key in (("Drawing numbers in more than one revision", "multi_revision"),
                       ("Older revisions that should probably be removed", "older"),
                       ("Extra copies of the same revision", "extra_copies"),
                       ("Drawings whose revisions cannot be ordered", "unordered"),
                       ("File names with no recognisable number/revision", "unparsed")):
        ws.append([label, stats[key]])


def write_revision_xlsx(rev_rows, stats, path, folders):
    """Workbook for --revisions-only runs (no content scan)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Folders scanned"])
    for f in folders:
        ws.append([str(f)])
    _add_revision_summary(ws, stats)
    ws.column_dimensions["A"].width = 60
    _add_revision_sheet(wb, rev_rows)
    wb.save(path)


def write_xlsx(results, path, folders, skipped, rev_rows=None, rev_stats=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {0: "E2F0D9", 1: "F8CBAD", 2: "FFE699", 3: "FFE699", 4: "FFF2CC", 5: "D9D9D9"}
    head_fill = PatternFill("solid", fgColor="1F3864")
    rows = [_row(r) for r in sorted(results, key=lambda r: (category_order(r["category"]), r["path"].lower()))]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Category", "Files"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
    for cat, n in summarize(results):
        ws.append([cat, n])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=fills[category_order(cat)[0]])
    ws.append([])
    st = Counter(r["status"] for r in results)
    ws.append(["Total PDFs scanned", len(results)])
    ws.append(["Flagged REVIEW", st.get(models.STATUS_REVIEW, 0)])
    ws.append(["Could not be read", st.get(models.STATUS_ERROR, 0)])
    if skipped:
        ws.append(["Non-PDF files ignored", sum(skipped.values())])
    ws.append([])
    ws.append(["Folders scanned"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for f in folders:
        ws.append([str(f)])
    if rev_stats is not None:
        _add_revision_summary(ws, rev_stats)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 12

    def sheet(title, subset):
        s = wb.create_sheet(title)
        s.append(COLUMNS)
        for c in s[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
        for row in subset:
            s.append([row[k] for k in COLUMNS])
            fill = PatternFill("solid", fgColor=fills[category_order(row["Category"])[0]])
            s.cell(s.max_row, 1).fill = fill
            ev = row["Evidence Image(s)"].split("; ")[0]
            if ev:
                cell = s.cell(s.max_row, COLUMNS.index("Evidence Image(s)") + 1)
                cell.hyperlink = "evidence/" + ev
                cell.font = Font(color="0563C1", underline="single")
        s.freeze_panes = "D2"
        s.auto_filter.ref = s.dimensions
        widths = {"Category": 38, "Status": 9, "File Name": 38, "Folder": 40, "ECCN (as written)": 18,
                  "Logo(s) Found": 14, "Logo Pages": 10, "Detection": 22, "Pages": 7, "Pages Scanned": 9,
                  "Size (KB)": 10, "Last Modified": 17, "Notes": 60, "Evidence Image(s)": 30, "Full Path": 60,
                  "Drawing Number": 20, "Revision": 9}
        for i, col in enumerate(COLUMNS, 1):
            s.column_dimensions[get_column_letter(i)].width = widths[col]
        for row in s.iter_rows(min_row=2):
            row[COLUMNS.index("Notes")].alignment = Alignment(wrap_text=True, vertical="top")

    sheet("All Files", rows)
    sheet("Needs Attention", [r for r in rows if r["Status"] != models.STATUS_OK])
    if rev_rows is not None:
        _add_revision_sheet(wb, rev_rows)
    wb.save(path)


def copy_sorted(results, dest):
    """Copy each PDF into <dest>/<category>[/<ECCN>]/<source folder name>/<relative path>.  Originals are untouched."""
    dest = Path(dest)
    n = 0
    for r in results:
        cat = r["category"]
        if cat.startswith(models.EAR_PREFIX):
            parts = ["EAR Controlled", cat[len(models.EAR_PREFIX):]]
        else:
            parts = [cat]
        root = Path(r["root"]) if r["root"] else Path(r["path"]).parent
        try:
            rel = Path(r["path"]).relative_to(root)
        except ValueError:
            rel = Path(Path(r["path"]).name)
        target = dest.joinpath(*[p.replace(os.sep, "_") for p in parts], root.name or "root", rel)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(r["path"], target)
        n += 1
    return n
